"""
Unit tests for template_builder.py
"""
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.utils import load_schema
from template_builder import (
    CUSTOM_FIELD,
    CUSTOM_HEADER_FILL,
    OPTIONAL_HEADER_FILL,
    REQUIRED_HEADER_FILL,
    TEMPLATE_SECTIONS,
    _section_fields,
    build_full_template,
    build_upload_template,
    load_field_descriptions,
)


def _rgb(fill):
    return fill.fgColor.rgb[-6:]


def _header(ws):
    return [cell for cell in ws[1] if cell.value is not None]


class TestBuildUploadTemplate:
    """Test cases for build_upload_template function."""

    @staticmethod
    def _template_header(template_bytes):
        return _header(load_workbook(BytesIO(template_bytes)).active)

    def test_required_then_optional_columns(self):
        header = self._template_header(build_upload_template(["a", "b"], ["c"]))
        assert [c.value for c in header] == ["a", "b", "c"]

    def test_required_and_optional_colours(self):
        header = self._template_header(build_upload_template(["a", "b"], ["c"]))
        assert _rgb(header[0].fill) == _rgb(REQUIRED_HEADER_FILL)
        assert _rgb(header[1].fill) == _rgb(REQUIRED_HEADER_FILL)
        assert _rgb(header[2].fill) == _rgb(OPTIONAL_HEADER_FILL)

    def test_no_optional_fields(self):
        for template in (
            build_upload_template(["a"], []),
            build_upload_template(["a"]),
        ):
            assert [c.value for c in self._template_header(template)] == ["a"]

    def test_duplicates_removed(self):
        header = self._template_header(build_upload_template(["a", "b"], ["b", "c"]))
        assert [c.value for c in header] == ["a", "b", "c"]

    def test_header_comment_includes_description(self):
        header = self._template_header(build_upload_template(["target_name"]))
        assert header[0].comment.text.startswith("Required. ")
        assert load_field_descriptions()["target_name"] in header[0].comment.text

    def test_template_reads_back_as_upload(self):
        template = build_upload_template(["target_name", "seq"], ["chrom"])
        df = pd.read_excel(BytesIO(template))
        assert df.columns.tolist() == ["target_name", "seq", "chrom"]


class TestBuildFullTemplate:
    """Test cases for build_full_template function."""

    def test_sheets(self):
        wb = load_workbook(BytesIO(build_full_template()))
        assert wb.sheetnames == ["Overview"] + [s["sheet"] for s in TEMPLATE_SECTIONS]

    def test_sheet_columns_match_schema(self):
        schema = load_schema()
        wb = load_workbook(BytesIO(build_full_template()))
        for section in TEMPLATE_SECTIONS:
            required, optional = _section_fields(schema, section["schema_path"])
            header = _header(wb[section["sheet"]])
            expected = list(dict.fromkeys(required + optional)) + [CUSTOM_FIELD]
            assert [c.value for c in header] == expected, section["sheet"]
            assert _rgb(header[-1].fill) == _rgb(CUSTOM_HEADER_FILL)

    def test_every_field_has_description(self):
        schema = load_schema()
        descriptions = load_field_descriptions()
        missing = {
            field
            for section in TEMPLATE_SECTIONS
            for field in sum(_section_fields(schema, section["schema_path"]), [])
            if field not in descriptions
        }
        assert not missing
