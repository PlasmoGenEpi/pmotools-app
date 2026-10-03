"""
Tests for full_template_loader.py: fill in the generated full template, load
it, merge it as the Create Final PMO page does, and validate the PMO.
"""
import datetime
import gzip
import json
import tarfile
import zipfile
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from pmotools.pmo_builder.merge_to_pmo import merge_to_pmo
from pmotools.pmo_engine.pmo_checker import PMOChecker, load_schema_by_version

from full_template_loader import (
    SECTION_BY_SHEET,
    bioinformatics_run_names,
    default_panel_name,
    load_full_template,
    read_full_template,
    read_template_upload,
    sheet_for_file_name,
)
from template_builder import build_full_template

FILLED_SHEETS = {
    "Panel Targets": [
        dict(
            target_name="t1",
            forward_primer_seq="ACGT",
            reverse_primer_seq="TTGA",
            reaction_name="pool1",
            chrom="chr1",
            insert_start=10,
            insert_end=50,
            strand="+",
            target_attributes="drug,diversity",
        ),
        dict(
            target_name="t2",
            forward_primer_seq="GGCC",
            reverse_primer_seq="AATT",
            reaction_name="pool1,pool2",
            chrom="chr2",
            insert_start=5,
            insert_end=60,
            strand="-",
        ),
    ],
    "Genome Info": [
        dict(
            name="Pf3D7",
            taxon_id=5833,
            genome_version="2020",
            url="https://plasmodb.org/Pf3D7.fasta",
            chromosomes="chr1,chr2",
        )
    ],
    "Microhaplotype": [
        dict(library_sample_name="L1", target_name="t1", seq="AAAA", reads=10),
        dict(library_sample_name="L1", target_name="t2", seq="CCCC", reads=5),
        dict(library_sample_name="L2", target_name="t1", seq="AAAT", reads=7),
    ],
    "Specimen Level": [
        dict(
            specimen_name="S1",
            collection_date=datetime.datetime(2024, 1, 2),
            collection_country="Kenya",
            project_name="proj1",
            host_subject_name="P1",
            storage_plate_name="plate1",
            storage_plate_row="a",
            storage_plate_col=1,
        ),
        dict(
            specimen_name="S2",
            collection_date="2024-01",
            collection_country="Kenya",
            project_name="proj1",
        ),
    ],
    "Project Info": [
        dict(
            project_name="proj1",
            project_description="A test project",
            project_contributors="Alice, Bob",
        )
    ],
    "Library Sample Info": [
        dict(
            library_sample_name="L1",
            specimen_name="S1",
            panel_name="panelA",
            sequencing_info_name="seq1",
        ),
        dict(
            library_sample_name="L2",
            specimen_name="S2",
            panel_name="panelA",
            sequencing_info_name="seq1",
        ),
    ],
    "Sequencing Info": [
        dict(
            sequencing_info_name="seq1",
            seq_platform="ILLUMINA",
            seq_instrument_model="MiSeq",
            library_layout="paired-end",
            library_strategy="AMPLICON",
            library_source="GENOMIC",
            library_selection="PCR",
            seq_date=datetime.datetime(2024, 2, 1),
        )
    ],
    "Bioinformatics Run": [
        dict(bioinformatics_run_name="run1", bioinformatics_method_name="pipeA"),
    ],
    "Bioinformatics Methods": [
        dict(
            program="cutadapt",
            program_version="4.0",
            bioinformatics_method_name="pipeA",
        ),
        dict(
            program="dada2",
            program_version="1.26",
            additional_argument="--maxEE 2, --truncQ 2",
            bioinformatics_method_name="pipeA",
        ),
        dict(
            program="other", program_version="1.0", bioinformatics_method_name="pipeB"
        ),
    ],
    "Raw Read Counts": [
        dict(library_sample_name="L1", total_raw_count=100),
        dict(library_sample_name="L2", total_raw_count=80),
    ],
    "Reads by Stage": [
        dict(
            library_sample_name="L1",
            target_name="t1",
            stage="demultiplexed",
            read_count=50,
        ),
        dict(
            library_sample_name="L2",
            target_name="t1",
            stage="demultiplexed",
            read_count=40,
        ),
    ],
}


def _filled_template(sheets=FILLED_SHEETS, rename=None):
    wb = load_workbook(BytesIO(build_full_template()))
    for sheet, rows in sheets.items():
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
        for old, new in (rename or {}).get(sheet, {}).items():
            ws.cell(row=1, column=header.index(old) + 1, value=new)
        for i, row in enumerate(rows, start=2):
            for field, value in row.items():
                ws.cell(row=i, column=header.index(field) + 1, value=value)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _merge(data):
    pmo = merge_to_pmo(
        mhap_info=data["microhaplotype_info"],
        panel_target_info=data["panel_info"],
        specimen_info=data.get("specimen_info"),
        library_sample_info=data.get("library_sample_info"),
        sequencing_info=data.get("seq_info"),
        bioinfo_method_info=data.get("bioinfo_methods_list"),
        bioinfo_run_info=data.get("bioinfo_run_infos"),
        project_info=data.get("project_info"),
        read_counts_by_stage_info=data.get("read_counts_per_stage"),
    )
    return json.loads(json.dumps(pmo, default=str))


def test_empty_sheets_skipped():
    sheets = read_full_template(
        _filled_template({"Microhaplotype": FILLED_SHEETS["Microhaplotype"]})
    )
    assert list(sheets) == ["Microhaplotype"]


def test_defaults_from_workbook():
    sheets = read_full_template(_filled_template())
    assert default_panel_name(sheets) == "panelA"
    assert bioinformatics_run_names(sheets) == ["run1"]


def test_full_template_builds_valid_pmo():
    sheets = read_full_template(_filled_template())
    data, errors = load_full_template(sheets, "panelA", "run1")
    assert errors == {}

    pmo = _merge(data)
    PMOChecker(load_schema_by_version("1.1.0")).validate_pmo_json(pmo)

    assert pmo["targeted_genomes"][0]["chromosomes"] == ["chr1", "chr2"]
    specimen = pmo["specimen_info"][0]
    assert specimen["collection_date"] == "2024-01-02"
    assert specimen["host_subject_name"] == "P1"
    assert specimen["storage_plate_info"]["plate_row"] == "A"
    assert pmo["project_info"][0]["project_contributors"] == ["Alice", "Bob"]
    assert pmo["sequencing_info"][0]["seq_date"] == "2024-02-01"
    assert [
        m["bioinformatics_method_name"] for m in pmo["bioinformatics_methods_info"]
    ] == ["pipeA", "pipeB"]
    assert pmo["bioinformatics_methods_info"][0]["methods"][1][
        "additional_argument"
    ] == ["--maxEE 2", "--truncQ 2"]
    assert pmo["bioinformatics_run_info"][0]["bioinformatics_methods_id"] == 0
    assert (
        len(pmo["read_counts_by_stage"][0]["read_counts_by_library_sample_by_stage"])
        == 2
    )


def test_alternate_column_names_recognised():
    rename = {
        "Panel Targets": {
            "forward_primer_seq": "fwd_primer",
            "target_name": "Target_ID",
        }
    }
    required = ("target_name", "forward_primer_seq", "reverse_primer_seq")
    rows = [{k: r[k] for k in required} for r in FILLED_SHEETS["Panel Targets"]]
    sheets = read_full_template(_filled_template({"Panel Targets": rows}, rename))
    data, errors = load_full_template(sheets, "panelA")
    assert errors == {}
    assert data["panel_info"]["target_info"][0]["forward_primer"]["seq"] == "ACGT"


def test_genome_id_accepts_genome_name():
    panel = [dict(r, genome_id="Pf3D7") for r in FILLED_SHEETS["Panel Targets"]]
    sheets = read_full_template(
        _filled_template(
            {"Panel Targets": panel, "Genome Info": FILLED_SHEETS["Genome Info"]}
        )
    )
    data, errors = load_full_template(sheets, "panelA")
    assert errors == {}
    assert data["panel_info"]["target_info"][0]["insert_location"]["genome_id"] == 0

    sheets["Panel Targets"]["genome_id"] = "unknown"
    _, errors = load_full_template(sheets, "panelA")
    assert "unknown" in errors["Panel Targets"]


def test_errors_reported_per_sheet():
    sheets = read_full_template(
        _filled_template(
            {
                "Panel Targets": FILLED_SHEETS["Panel Targets"],
                "Microhaplotype": FILLED_SHEETS["Microhaplotype"],
                "Bioinformatics Run": [
                    dict(
                        bioinformatics_run_name="run1",
                        bioinformatics_method_name="missing",
                    )
                ],
                "Bioinformatics Methods": FILLED_SHEETS["Bioinformatics Methods"],
            }
        )
    )
    data, errors = load_full_template(
        sheets, panel_name="", bioinformatics_run_name="run1"
    )
    assert "panel name" in errors["Panel Targets"]
    assert "not on the Bioinformatics Methods sheet" in errors["Bioinformatics Run"]
    assert "microhaplotype_info" in data
    assert "bioinfo_methods_list" in data


class _Upload(BytesIO):
    def __init__(self, name, data):
        super().__init__(data)
        self.name = name


def _table_bytes(sheet, sep=","):
    df = pd.DataFrame(FILLED_SHEETS[sheet])
    df = df.map(
        lambda v: v.strftime("%Y-%m-%d") if isinstance(v, datetime.datetime) else v
    )
    return df.to_csv(sep=sep, index=False).encode()


def _assert_valid_pmo(sheets):
    data, errors = load_full_template(sheets, "panelA", "run1")
    assert errors == {}
    pmo = _merge(data)
    PMOChecker(load_schema_by_version("1.1.0")).validate_pmo_json(pmo)
    return pmo


def test_sheet_for_file_name():
    assert sheet_for_file_name("Panel Targets.csv") == "Panel Targets"
    assert sheet_for_file_name("panel_targets.tsv") == "Panel Targets"
    assert sheet_for_file_name("folder/panel_info_template.csv") == "Panel Targets"
    assert sheet_for_file_name("specimen-level.csv.gz") == "Specimen Level"
    assert sheet_for_file_name("raw_counts.txt") == "Raw Read Counts"
    assert sheet_for_file_name("notes.csv") is None


def test_csv_and_tsv_files_build_valid_pmo():
    files = [
        _Upload(
            f"{sheet.lower().replace(' ', '_')}.{'tsv' if i % 2 else 'csv'}",
            _table_bytes(sheet, "\t" if i % 2 else ","),
        )
        for i, sheet in enumerate(FILLED_SHEETS)
    ]
    files.append(
        _Upload("specimen_level.csv.gz", gzip.compress(_table_bytes("Specimen Level")))
    )
    sheets, ignored = read_template_upload(files)
    assert ignored == []
    assert set(sheets) == set(FILLED_SHEETS)
    pmo = _assert_valid_pmo(sheets)
    assert pmo["specimen_info"][0]["collection_date"] == "2024-01-02"


def test_zip_folder_builds_valid_pmo():
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for sheet in FILLED_SHEETS:
            section = SECTION_BY_SHEET[sheet]
            archive.writestr(f"pmo/{section['template_name']}.csv", _table_bytes(sheet))
        archive.writestr("pmo/README.md", "notes")
        archive.writestr("__MACOSX/pmo/._panel_info.csv", "junk")
        archive.writestr("pmo/.DS_Store", "junk")
    sheets, ignored = read_template_upload([_Upload("pmo.zip", buffer.getvalue())])
    assert ignored == ["pmo/README.md"]
    assert set(sheets) == set(FILLED_SHEETS)
    _assert_valid_pmo(sheets)


def test_tar_gz_folder():
    buffer = BytesIO()
    with tarfile.open(fileobj=buffer, mode="w:gz") as archive:
        for sheet in ("Bioinformatics Run", "Microhaplotype"):
            data = _table_bytes(sheet, "\t")
            info = tarfile.TarInfo(f"pmo/{sheet}.tsv")
            info.size = len(data)
            archive.addfile(info, BytesIO(data))
    sheets, ignored = read_template_upload([_Upload("pmo.tar.gz", buffer.getvalue())])
    assert ignored == []
    assert list(sheets["Bioinformatics Run"].columns) == [
        "bioinformatics_run_name",
        "bioinformatics_method_name",
    ]
    assert len(sheets["Microhaplotype"]) == 3


def test_excel_upload():
    workbook = _filled_template().getvalue()
    sheets, ignored = read_template_upload(
        [_Upload("PMO_building_template.xlsx", workbook)]
    )
    assert ignored == []
    assert set(sheets) == set(FILLED_SHEETS)


def test_minimum_pmo_example_data():
    folder = Path(__file__).parent.parent / "example_data" / "minimum_pmo"
    for names in (
        ["panel_info.tsv", "microhaplotype_info.tsv"],
        ["minimum_pmo_template.xlsx"],
    ):
        files = [_Upload(n, (folder / n).read_bytes()) for n in names]
        sheets, ignored = read_template_upload(files)
        assert ignored == []
        assert set(sheets) == {"Panel Targets", "Microhaplotype"}
        data, errors = load_full_template(sheets, "Mad4hatter")
        assert errors == {}
        pmo = json.loads(
            json.dumps(
                merge_to_pmo(
                    mhap_info=data["microhaplotype_info"],
                    panel_target_info=data["panel_info"],
                ),
                default=str,
            )
        )
        PMOChecker(load_schema_by_version("1.1.0")).validate_pmo_json(pmo)


def test_missing_required_column():
    rows = [
        {k: v for k, v in r.items() if k != "reads"}
        for r in FILLED_SHEETS["Microhaplotype"]
    ]
    sheets = read_full_template(_filled_template({"Microhaplotype": rows}))
    sheets["Microhaplotype"] = sheets["Microhaplotype"].drop(
        columns=["reads"], errors="ignore"
    )
    _, errors = load_full_template(sheets)
    assert "reads" in errors["Microhaplotype"]
