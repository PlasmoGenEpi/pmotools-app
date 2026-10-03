"""
Excel templates for the tables uploaded to the PMO Builder.

Column names come from conf/schema.json, which mirrors the column arguments of
the pmotools-python table-to-PMO builders, so a filled-in template can be
uploaded to the app or passed straight to pmotools-python.
"""
import json
from io import BytesIO

import streamlit as st
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.utils import load_schema

FIELD_DESCRIPTIONS_PATH = "conf/field_descriptions.json"
CUSTOM_FIELD = "custom_field"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

REQUIRED_HEADER_FILL = PatternFill("solid", fgColor="DDEE3B")
OPTIONAL_HEADER_FILL = PatternFill("solid", fgColor="FCDAE6")
CUSTOM_HEADER_FILL = PatternFill("solid", fgColor="7589C3")

TEMPLATE_SECTIONS = [
    {
        "sheet": "Panel Targets",
        "schema_path": ["panel_info"],
        "page": "Panel Information",
        "description": "One row per target in the panel, with its primers and optional genomic locations. The panel name is entered in the app.",
        "pmo_sections": "panel_info, target_info",
    },
    {
        "sheet": "Genome Info",
        "schema_path": ["targeted_genomes"],
        "page": "Panel Information",
        "description": "The genome(s) the panel targets. Required if Panel Targets includes genomic locations.",
        "pmo_sections": "targeted_genomes",
    },
    {
        "sheet": "Microhaplotype",
        "schema_path": ["mhap_info"],
        "page": "Microhaplotype Information",
        "description": "One row per microhaplotype detected in each library sample for each target. The bioinformatics run name is entered in the app.",
        "pmo_sections": "representative_microhaplotypes, detected_microhaplotypes",
    },
    {
        "sheet": "Specimen Level",
        "schema_path": ["specimen_level_metadata"],
        "page": "Specimen Level Metadata",
        "description": "Metadata on each biological specimen.",
        "pmo_sections": "specimen_info",
    },
    {
        "sheet": "Project Info",
        "schema_path": ["project_name"],
        "page": "Project Information",
        "description": "The project(s) the specimens belong to.",
        "pmo_sections": "project_info",
    },
    {
        "sheet": "Library Sample Info",
        "schema_path": ["library_sample_level_metadata"],
        "page": "Library Sample Level Metadata",
        "description": "Each library prepared from a specimen; a specimen may have several library samples.",
        "pmo_sections": "library_sample_info",
    },
    {
        "sheet": "Sequencing Info",
        "schema_path": ["sequencing_info"],
        "page": "Sequencing Information",
        "description": "How the samples were sequenced, e.g. platform, instrument and dates.",
        "pmo_sections": "sequencing_info",
    },
    {
        "sheet": "Bioinformatics Run",
        "schema_path": ["bioinformatics_run_info"],
        "page": "Bioinformatics Information",
        "description": "Each bioinformatics run, if samples were analysed in one or several runs.",
        "pmo_sections": "bioinformatics_run_info",
    },
    {
        "sheet": "Bioinformatics Methods",
        "schema_path": ["bioinformatics_methods_info"],
        "page": "Bioinformatics Information",
        "description": "The programs used to generate the microhaplotypes.",
        "pmo_sections": "bioinformatics_methods_info",
    },
    {
        "sheet": "Raw Read Counts",
        "schema_path": ["read_counts_perstage", "raw_counts"],
        "page": "Read Counts per Stage",
        "description": "Total raw read count for each library sample.",
        "pmo_sections": "read_counts_by_stage",
    },
    {
        "sheet": "Reads by Stage",
        "schema_path": ["read_counts_perstage", "reads_by_stage"],
        "page": "Read Counts per Stage",
        "description": "Read counts per library sample, per target, at each pipeline stage.",
        "pmo_sections": "read_counts_by_stage",
    },
]

OVERVIEW_TEXT = [
    "This template provides a structured format for building a Portable Microhaplotype Object (PMO).",
    "Find more details at: https://plasmogenepi.github.io/PMO_Docs/",
    "Each sheet is one table. Required columns must be completed to create a valid PMO; optional columns store additional metadata. Hover over a column header to see a description of the field.",
    "Column names match the PMO Builder app and the table-to-PMO functions in pmotools-python, so a completed sheet can be uploaded to the matching page of the app (select the sheet after uploading this workbook) or passed directly to pmotools-python.",
    "Some values are entered in the app rather than in a sheet: the panel name (Panel Information page) and the bioinformatics run name (Microhaplotype Information and Read Counts per Stage pages).",
    "Columns that accept several values (noted in the header description) should be separated with commas. Dates should be stored as YYYY-MM-DD.",
    "You may add your own custom fields to store project-specific information. If there are fields you use regularly and believe would be broadly useful, please contact us so we can consider them for a future version of the PMO specification.",
]


def load_field_descriptions() -> dict:
    with open(FIELD_DESCRIPTIONS_PATH, "r") as f:
        return json.load(f)


def _section_fields(schema: dict, schema_path: list) -> tuple:
    section = schema
    for key in schema_path:
        section = section[key]
    return section["required"], section.get("optional", [])


def _write_header(ws, required_fields, optional_fields, descriptions, custom=False):
    required = list(dict.fromkeys(required_fields))
    optional = [f for f in dict.fromkeys(optional_fields or []) if f not in required]
    columns = [(f, "Required", REQUIRED_HEADER_FILL) for f in required]
    columns += [(f, "Optional", OPTIONAL_HEADER_FILL) for f in optional]
    if custom:
        columns.append((CUSTOM_FIELD, "Custom", CUSTOM_HEADER_FILL))

    for col, (field, status, fill) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=field)
        cell.fill = fill
        description = descriptions.get(field)
        note = f"{status}. {description}" if description else status
        comment = Comment(note, "PMO Builder")
        comment.width, comment.height = 300, 120
        cell.comment = comment
        ws.column_dimensions[get_column_letter(col)].width = max(12, len(field) + 2)
    ws.freeze_panes = "A2"


def _to_bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_upload_template(required_fields, optional_fields=None) -> bytes:
    """Return a single-sheet Excel template whose header row lists required then
    optional fields, coloured to match the full PMO building template."""
    wb = Workbook()
    ws = wb.active
    ws.title = "template"
    _write_header(ws, required_fields, optional_fields, load_field_descriptions())
    return _to_bytes(wb)


def _write_overview(ws):
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 40

    ws.cell(row=1, column=1, value="Overview").font = Font(bold=True, size=14)
    row = 2
    for text in OVERVIEW_TEXT:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = ws.cell(row=row, column=1, value=text)
        cell.alignment = wrap
        ws.row_dimensions[row].height = 15 * (1 + len(text) // 140)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Colour Key").font = bold
    for label, fill in [
        ("Required in format", REQUIRED_HEADER_FILL),
        ("Not required in format", OPTIONAL_HEADER_FILL),
        (
            "Not defined by schema (custom fields are allowed, but not recommended)",
            CUSTOM_HEADER_FILL,
        ),
    ]:
        row += 1
        ws.cell(row=row, column=1).fill = fill
        ws.cell(row=row, column=2, value=label)

    row += 2
    for col, heading in enumerate(
        ["Sheet", "App page", "Description", "Associated PMO sections"], start=1
    ):
        ws.cell(row=row, column=col, value=heading).font = bold
    for section in TEMPLATE_SECTIONS:
        row += 1
        values = [
            section["sheet"],
            section["page"],
            section["description"],
            section["pmo_sections"],
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value).alignment = wrap


def build_full_template() -> bytes:
    """Return the full PMO building template: an Overview sheet plus one sheet
    per uploadable table."""
    schema = load_schema()
    descriptions = load_field_descriptions()
    wb = Workbook()
    _write_overview(wb.active)
    wb.active.title = "Overview"
    for section in TEMPLATE_SECTIONS:
        required, optional = _section_fields(schema, section["schema_path"])
        ws = wb.create_sheet(section["sheet"])
        _write_header(ws, required, optional, descriptions, custom=True)
    return _to_bytes(wb)


def render_template_download(
    required_fields, optional_fields=None, template_name="upload", key_suffix=""
):
    """Render a download button for an Excel template of the expected upload columns."""
    st.download_button(
        label="Download template",
        data=build_upload_template(required_fields, optional_fields),
        file_name=f"{template_name}_template.xlsx",
        mime=XLSX_MIME,
        icon=":material/download:",
        key=f"template_download_{key_suffix or template_name}",
    )
